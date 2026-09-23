"""
regional-indicators/fetchers/fetch_business.py
-----------------------------------------------------------------------
Fetches ABS business counts by SA2, industry division, and turnover
size range -- "Counts of Australian Businesses, including Entries and
Exits", Data cube 9.

Source (confirmed real, downloaded and inspected directly, 2026-09-23):
  https://www.abs.gov.au/statistics/economy/business-indicators/counts-australian-businesses-including-entries-and-exits/{release}/8165DC09.xlsx
Current release used: jul2021-jun2025 (published 16 Dec 2025).

REAL STRUCTURE, CONFIRMED DIRECTLY (not assumed from documentation
alone) -- one table PER YEAR, not one combined time series:
  Table 1 = June 2025 (the latest year)
  Table 2 = June 2024
  Table 3 = June 2023
This fetcher only uses Table 1 (the latest year) -- the workbook
already has history for earlier years from past updates; this run
only needs to add whatever new year Table 1 represents.

REAL COLUMNS (row 6-7 headers, data from row 8), confirmed directly:
  A: Industry Code    B: Industry Label
  C: SA2 Code          D: SA2 Label
  E: 0-<$50k no.        F: $50k-<$200k no.
  G: $200k-<$2m no.     H: $2m-<$5m no.
  I: $5m-<$10m no.      J: $10m+ no.
  K: Total no.

REAL DATA HAZARD, CONFIRMED DIRECTLY: every industry code appears
TWICE -- once as real per-SA2 detail rows, once as a "Total <industry>"
row with SA2 Code and SA2 Label both blank (a state/national aggregate,
not per-SA2 data). Filtered out by requiring a non-blank SA2 code --
trivial once you know to check for it, silently wrong (double-counted
totals) if you don't.

METHODOLOGY, per Notes_DD.docx (Steve's documented process) and
confirmed against the real data:
  - Industry code "A" (Agriculture, Forestry and Fishing) = Primary
    Production. ALL other codes (including "X", Currently Unknown) =
    Non-Primary Production, summed together -- the notes are explicit
    ("group all OTHER industries as non-primary production"), so X is
    deliberately included in NPP, not excluded as ambiguous.
  - Turnover bands: keep the first three as published (0-50k, 50k-200k,
    200k-2m); combine 2m-5m + 5m-10m + 10m+ into a single "2m+" category
    -- matches the workbook's four bands exactly.
  - SA2 rows are summed within each production category for a single
    SA2; a composite area (Toowoomba) sums across ALL its component
    SA2s for each turnover band.

SA2 MAPPING, confirmed 2026-09-23 directly against the real Business
sheet and towns.toml, with corrections found along the way:
  Chinchilla=307011172  Goondiwindi=307011173  Moranbah=312011341
  Tara=307011178  Wambo(Dalby)=307021183  Miles-Wandoan=307011175
  Roma=307011176  Roma Surrounds(Wallumbilla)=307011177
  Broadsound-Nebo(Dysart)=312011338  Narrabri=110031197
  Narrabri Surrounds=110031198 (no corresponding project town)
  Toowoomba - SA2 Composite = 317011456 + 317011454 + 317011458
    (+ 317011457, "Toowoomba - East", confirmed to exist via live ABS
    lookup but not yet in towns.toml -- see TODO.md, this is a real
    open decision, not resolved by this fetcher unilaterally. Currently
    NOT included in the composite sum below pending that decision.)
"""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from fetchers.base import BaseFetcher

try:
    import openpyxl
    import requests
except ImportError:
    raise ImportError("pip install openpyxl requests")


ABS_RELEASE = "jul2021-jun2025"
ABS_DC9_URL = (
    f"https://www.abs.gov.au/statistics/economy/business-indicators/"
    f"counts-australian-businesses-including-entries-and-exits/"
    f"{ABS_RELEASE}/8165DC09.xlsx"
)
PRIMARY_PRODUCTION_CODE = "A"

# Confirmed real SA2 codes for every region the Business sheet tracks,
# 2026-09-23 (see module docstring for the verification trail).
SA2_REGIONS = {
    "Broadsound-Nebo":            [312011338],
    "Chinchilla":                 [307011172],
    "Goondiwindi":                [307011173],
    "Miles-Wandoan":              [307011175],
    "Moranbah":                   [312011341],
    "Narrabri":                   [110031197],
    "Narrabri Surrounds":         [110031198],
    "Roma":                       [307011176],
    "Roma Surrounds":             [307011177],
    "Tara":                       [307011178],
    "Toowoomba - SA2 Composite":  [317011456, 317011454, 317011458],
    "Wambo":                      [307021183],
}


class ABSBusinessFetcher(BaseFetcher):

    SOURCE_NAME      = "business"
    SUPPORTED_STATES = []

    def fetch_all(self):
        cache_key = f"abs_dc9_{ABS_RELEASE}"
        path = self.download(ABS_DC9_URL, cache_key, suffix=".xlsx")
        if not path:
            self.result.add_error("ALL", "Data cube 9 download failed")
            return

        self.log.info(f"  Parsing {path.name} ({path.stat().st_size // 1024} KB)")

        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            ws = wb["Table 1"]  # confirmed: the latest year, see module docstring
        except Exception as exc:
            self.log.error(f"Could not open Table 1: {exc}", exc_info=True)
            self.result.add_error("ALL", f"Could not open Table 1: {exc}")
            return

        # year label, confirmed real location: row 4, e.g.
        # "Businesses by Industry Division by SA2 by Turnover Size Ranges, June 2025 (a) (b)"
        title_cell = ws.cell(4, 1).value or ""
        self.log.info(f"  Table 1 title: {title_cell}")

        by_sa2 = self._parse_table(ws)
        if not by_sa2:
            self.result.add_error("ALL", "No data rows parsed from Table 1")
            return

        for region_name, sa2_codes in SA2_REGIONS.items():
            self._write_region(region_name, sa2_codes, by_sa2)

    def _parse_table(self, ws) -> dict:
        """Returns {sa2_code: {"NPP": {band: count}, "PP": {band: count}}}
        aggregated straight from the raw rows -- turnover bands combined
        into the workbook's 4 categories here, industry rows split into
        NPP/PP here, so downstream region-writing is a simple sum.
        """
        by_sa2: dict = defaultdict(lambda: {
            "NPP": defaultdict(int), "PP": defaultdict(int),
        })
        rows_used, rows_skipped_total, rows_skipped_other = 0, 0, 0

        for row in ws.iter_rows(min_row=8, values_only=True):
            industry_code = row[0]
            if industry_code is None:
                break  # end of data block

            sa2_code = row[2]
            if sa2_code is None:
                rows_skipped_total += 1
                continue  # a "Total <industry>" aggregate row, not per-SA2

            try:
                sa2_code = int(sa2_code)
            except (TypeError, ValueError):
                rows_skipped_other += 1
                continue

            band = "PP" if industry_code == PRIMARY_PRODUCTION_CODE else "NPP"
            b_0_50k, b_50_200k, b_200k_2m, b_2_5m, b_5_10m, b_10m_plus = (
                row[4] or 0, row[5] or 0, row[6] or 0, row[7] or 0, row[8] or 0, row[9] or 0
            )

            entry = by_sa2[sa2_code][band]
            entry["0k-50k"]   += b_0_50k
            entry["50k-200k"] += b_50_200k
            entry["200k-2m"]  += b_200k_2m
            entry["2m+"]      += (b_2_5m + b_5_10m + b_10m_plus)
            rows_used += 1

        self.log.info(
            f"  Parsed {rows_used} rows ({len(by_sa2)} SA2s), "
            f"skipped {rows_skipped_total} state/national totals, "
            f"{rows_skipped_other} unparseable"
        )
        return dict(by_sa2)

    def _write_region(self, region_name: str, sa2_codes: list, by_sa2: dict):
        combined = {"NPP": defaultdict(int), "PP": defaultdict(int)}
        found_any = False

        for sa2_code in sa2_codes:
            if sa2_code not in by_sa2:
                self.log.warning(f"  [{region_name}] SA2 {sa2_code} not found in Table 1")
                continue
            found_any = True
            for production in ("NPP", "PP"):
                for band, count in by_sa2[sa2_code][production].items():
                    combined[production][band] += count

        if not found_any:
            self.log.warning(f"  [{region_name}] no SA2 data found at all -- skipping")
            self.result.towns_failed.append(region_name)
            return

        out = {
            "region": region_name,
            "sa2_codes": sa2_codes,
            "source": "ABS Counts of Australian Businesses, Data cube 9",
            "source_url": ABS_DC9_URL,
            "note": (
                f"Industry code 'A' (Agriculture, Forestry and Fishing) = Primary "
                f"Production; all other industry codes (including 'X', Currently "
                f"Unknown) = Non-Primary Production, summed. Turnover bands "
                f"2m-5m/5m-10m/10m+ combined into a single '2m+' category. "
                f"{'Composite of ' + str(len(sa2_codes)) + ' SA2s.' if len(sa2_codes) > 1 else ''}"
            ),
            "indicators": {
                "npp": {band: count for band, count in combined["NPP"].items()},
                "pp":  {band: count for band, count in combined["PP"].items()},
            },
        }

        out_dir = Path(__file__).parent.parent / "cache" / "business"
        out_dir.mkdir(parents=True, exist_ok=True)
        slug = region_name.lower().replace(" ", "_").replace("-", "_")
        out_path = out_dir / f"{slug}_business.json"
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(out, f, indent=2)

        self.log.info(
            f"  {region_name}: NPP total={sum(combined['NPP'].values())}, "
            f"PP total={sum(combined['PP'].values())} -> {out_path.name}"
        )
        self.result.towns_ok.append(region_name)


if __name__ == "__main__":
    from logger import get_logger
    get_logger()
    result = ABSBusinessFetcher().run()
    sys.exit(0 if result.success else 1)