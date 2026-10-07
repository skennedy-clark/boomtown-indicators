"""
regional-indicators/fetchers/fetch_narrabri_sales_rent.py

Fetches mean and median sales price, sales number, and median weekly
rent for new bonds (House 3-bedroom, and Total 3-bedroom) for Narrabri
(LGA), for the Narrabri block of the Housing sheet (rows 132-141).

Source: NSW Department of Communities and Justice (DCJ) Rent and Sales
Report: quarterly, LGA level, since 2011. Its "New Bonds" terminology
matches the sheet's row labels.
  https://dcj.nsw.gov.au/about-us/families-and-communities-statistics/housing-rent-and-sales/rent-and-sales-report.html

File layout:
  - Each quarter has two files, a "Sales tables" workbook and a "Rent
    tables" workbook, each with an "LGA" sheet.
  - Sales LGA sheet: long format, one row per (LGA, DwellingType).
    Columns "Median Sales Price $'000s", "Mean Sales Price $'000s" and
    "Sales No.". Prices are in thousands and are multiplied by 1000 to
    match the sheet, which holds whole dollars (282034.72, not 282.03).
    Narrabri's row is the one with DwellingType="Total" (all dwelling
    types); for example Q1 2026 gives median $435k, mean $458k and 51
    sales.
  - Rent LGA sheet: long format, one row per (LGA, DwellingType,
    Bedrooms). Column "Median Weekly Rent for New Bonds $". Two rows
    are read: DwellingType="House" with Bedrooms="3 Bedrooms" (the
    sheet's "House 3-Bed" row), and DwellingType="Total" with
    Bedrooms="3 Bedrooms" (the sheet's "Total" row, which is the
    all-types 3-bedroom figure, not an all-bedrooms figure).

File naming: file names are inconsistent between quarters, including
within one calendar year ("issue-152-sales-tables-mar-2025.xlsx"
against "sales_tables_june_2025_quarter.xlsx": underscores or hyphens,
with or without an issue number). Historical file names therefore
cannot be derived, and are handled as follows:
  - The file names for 2025 (four quarters, sales and rent) are listed
    in QUARTER_FILES as a backfill.
  - The current quarter's files are discovered on each run by scraping
    the "rent-and-sales-report.html" landing page, which links the
    latest quarter as "{type}-tables-{month}-{year}-quarter.xlsx". Rent
    and sales are not published for the same quarter, because their
    reporting lags differ (for example Rent = June 2026 with Sales =
    March 2026).
  - Update each cycle: once all four quarters of a new calendar year
    exist, add them to QUARTER_FILES, copying the file names from the
    "Previous rent and sales reports" archive page.

Aggregation: as for region-level Housing data in fetch_qgso_housing.py
(_fetch_regions). Price and rent are the mean of the four quarterly
values for the year; Sales No. is the sum of the four quarterly counts.
Complete years only (all four quarters present).

Output: cache/housing/regions/narrabri_sales_rent.json.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from fetchers.base import BaseFetcher
from config import CACHE_DIR, YEAR_START, YEAR_END

try:
    import requests
    import openpyxl
except ImportError:
    raise ImportError("pip install requests openpyxl")


DAM_BASE = (
    "https://dcj.nsw.gov.au/content/dam/dcj/dcj-website/documents/"
    "about-us/families-and-communities-statistics/housing-and-rent-sales"
)
LANDING_PAGE = (
    "https://dcj.nsw.gov.au/about-us/families-and-communities-statistics/"
    "housing-rent-and-sales/rent-and-sales-report.html"
)

NARRABRI_NAME = "Narrabri"

# Update each cycle. Historical file names, each of which resolves on
# the DCJ site. Add a year's quarters once the full calendar year is
# available; the names cannot be discovered automatically (see the
# module docstring).
QUARTER_FILES = {
    # (year, quarter): (sales_filename, rent_filename)
    (2025, 1): ("issue-152-sales-tables-mar-2025.xlsx", "issue-151-rent-tables-mar-2025.xlsx"),
    (2025, 2): ("sales_tables_june_2025_quarter.xlsx", "Rent_tables_June_2025_quarter.xlsx"),
    (2025, 3): ("sales_tables_september_2025_quarter.xlsx", "rent_tables_september_2025_quarter.xlsx"),
    (2025, 4): ("sales-tables-december-2025-quarter.xlsx", "rent_tables_december_2025_quarter.xlsx"),
}

SALES_LABELS = {
    "mean_sales_price":   "Mean Sales Price $",
    "median_sales_price": "Median Sales Price $",
    "sales_no":           "Sales No.",
}
RENT_LABELS = {
    "rent_house_3bed":  "Median Weekly Rent for New Bonds $ (House, 3 Bedrooms)",
    "rent_total_3bed":  "Median Weekly Rent for New Bonds $ (Total, 3 Bedrooms)",
}


def _discover_current_quarter_files(log) -> dict[str, str] | None:
    """Return the current quarter's file URLs, scraped from the landing
    page, as {"sales": url, "rent": url}.

    The files are linked as /content/dam/... URLs under button text of
    the form "Rent tables <Month> <Year> quarter" and "Sales tables
    <Month> <Year> quarter". Returns None if both links cannot be found,
    in which case only the QUARTER_FILES backfill is used.
    """
    try:
        resp = requests.get(LANDING_PAGE, timeout=30)
        resp.raise_for_status()
    except requests.RequestException as exc:
        log.warning(f"  Could not fetch the landing page: {exc}")
        return None

    found = {}
    for kind in ("rent", "sales"):
        m = re.search(
            rf'/content/dam/dcj/dcj-website/documents/about-us/families-and-communities-statistics/'
            rf'housing-and-rent-sales/({kind}-tables-[a-z]+-\d{{4}}-quarter\.xlsx)',
            resp.text,
            re.IGNORECASE,
        )
        if m:
            found[kind] = f"{DAM_BASE}/{m.group(1)}"
    if len(found) == 2:
        log.info(f"  Discovered current quarter files live: {found}")
        return found
    log.warning("  Could not find both current-quarter links on the landing page")
    return None


def _month_to_quarter(month_name: str) -> int:
    m = month_name.lower()
    if m in ("march", "mar"): return 1
    if m in ("june", "jun"): return 2
    if m in ("september", "sep"): return 3
    if m in ("december", "dec"): return 4
    return 0


def _parse_sales_file(path: Path, log) -> dict | None:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["LGA"]
    except Exception as exc:
        log.warning(f"  Could not open/parse {path.name}: {exc}")
        return None
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
        if row[3] == "Local Government Area (LGA)":
            header_row = row
            break
    if not header_row:
        log.warning(f"  Could not find the header row in {path.name}")
        return None
    cols = {v: i for i, v in enumerate(header_row) if v}
    median_col = cols.get("Median Sales Price\n$'000s")
    mean_col = cols.get("Mean Sales Price\n$'000s")
    no_col = cols.get("Sales\nNo.")
    if median_col is None or mean_col is None or no_col is None:
        log.warning(f"  Column headers didn't match expected names in {path.name}")
        return None

    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[3] == NARRABRI_NAME and row[4] == "Total":
            def _num(v):
                if v in (None, "-", "s"):
                    return None
                try:
                    return float(str(v).replace(",", ""))
                except ValueError:
                    return None
            median, mean, no = _num(row[median_col]), _num(row[mean_col]), _num(row[no_col])
            if median is None or mean is None or no is None:
                return None
            return {"median_sales_price": median * 1000, "mean_sales_price": mean * 1000, "sales_no": no}
    return None


def _parse_rent_file(path: Path, log) -> dict | None:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["LGA"]
    except Exception as exc:
        log.warning(f"  Could not open/parse {path.name}: {exc}")
        return None
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
        if row[3] == "Local Government Area (LGA)":
            header_row = row
            break
    if not header_row:
        log.warning(f"  Could not find the header row in {path.name}")
        return None
    cols = {v: i for i, v in enumerate(header_row) if v}
    median_col = cols.get("Median Weekly Rent for New Bonds\n$")
    if median_col is None:
        log.warning(f"  Rent column header didn't match expected name in {path.name}")
        return None

    def _num(v):
        if v in (None, "-", "s"):
            return None
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            return None

    out = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[3] != NARRABRI_NAME or row[5] != "3 Bedrooms":
            continue
        if row[4] == "House":
            out["rent_house_3bed"] = _num(row[median_col])
        elif row[4] == "Total":
            out["rent_total_3bed"] = _num(row[median_col])
    return out if out else None


class NarrabriSalesRentFetcher(BaseFetcher):

    SOURCE_NAME      = "narrabri_sales_rent"
    SUPPORTED_STATES = ["NSW"]

    def fetch_all(self):
        quarters = dict(QUARTER_FILES)  # (year, quarter) -> (sales_fn, rent_fn)

        current = _discover_current_quarter_files(self.log)
        # The sales and rent files on the landing page are not necessarily
        # for the same quarter (rent lags less than sales). Each is filed
        # under its own reporting quarter, parsed from its file name.
        current_urls: dict[tuple[int, int], dict[str, str]] = {}
        if current:
            for kind, url in current.items():
                m = re.search(r'-([a-z]+)-(\d{4})-quarter\.xlsx', url, re.IGNORECASE)
                if m:
                    q = _month_to_quarter(m.group(1))
                    yr = int(m.group(2))
                    if q:
                        current_urls.setdefault((yr, q), {})[kind] = url

        by_quarter_data: dict[tuple[int, int], dict] = {}

        for (yr, q), (sales_fn, rent_fn) in quarters.items():
            sales_url = f"{DAM_BASE}/{sales_fn}"
            rent_url = f"{DAM_BASE}/{rent_fn}"
            sales_path = self.download(sales_url, f"dcj_sales_{yr}q{q}", suffix=".xlsx")
            rent_path = self.download(rent_url, f"dcj_rent_{yr}q{q}", suffix=".xlsx")
            sales_data = _parse_sales_file(sales_path, self.log) if sales_path else None
            rent_data = _parse_rent_file(rent_path, self.log) if rent_path else None
            if sales_data or rent_data:
                merged = {}
                merged.update(sales_data or {})
                merged.update(rent_data or {})
                by_quarter_data[(yr, q)] = merged
                self.log.info(f"  [{yr} Q{q}] parsed: {merged}")

        for (yr, q), urls in current_urls.items():
            if (yr, q) in by_quarter_data:
                continue  # already loaded from QUARTER_FILES
            merged = {}
            if "sales" in urls:
                p = self.download(urls["sales"], f"dcj_sales_{yr}q{q}", suffix=".xlsx")
                if p:
                    merged.update(_parse_sales_file(p, self.log) or {})
            if "rent" in urls:
                p = self.download(urls["rent"], f"dcj_rent_{yr}q{q}", suffix=".xlsx")
                if p:
                    merged.update(_parse_rent_file(p, self.log) or {})
            if merged:
                by_quarter_data[(yr, q)] = merged
                self.log.info(f"  [{yr} Q{q}] (live-discovered) parsed: {merged}")

        if not by_quarter_data:
            self.result.add_error("ALL", "No Narrabri sales/rent data assembled from any quarter")
            return

        # Aggregate by calendar year: mean for price and rent, sum for the
        # sales count.
        by_year: dict[int, dict[str, list]] = {}
        for (yr, q), vals in by_quarter_data.items():
            if not (YEAR_START <= yr <= YEAR_END):
                continue
            for key, v in vals.items():
                if v is not None:
                    by_year.setdefault(yr, {}).setdefault(key, []).append(v)

        all_keys = list(SALES_LABELS) + list(RENT_LABELS)
        indicators = {}
        for key in all_keys:
            label = SALES_LABELS.get(key) or RENT_LABELS.get(key)
            values = {}
            for yr, series in by_year.items():
                vals = series.get(key, [])
                if len(vals) != 4:
                    continue  # complete years (4 quarters) only
                if key == "sales_no":
                    values[str(yr)] = int(sum(vals))
                else:
                    values[str(yr)] = round(sum(vals) / len(vals), 2)
            indicators[key] = {"label": label, "values": values}

        out = {
            "region":     NARRABRI_NAME,
            "section":    "Narrabri",
            "source":     "NSW Department of Communities and Justice (DCJ), Rent and Sales Report — LGA tables",
            "source_url": LANDING_PAGE,
            "note": (
                "Calendar-year means of the 4 quarterly Median/Mean Sales Price and "
                "Median Weekly Rent readings, and SUM of the 4 quarterly Sales No. "
                "counts -- complete years only (all 4 quarters present). Sales Price "
                "converted from the source's $'000s to whole dollars to match this "
                "sheet's existing convention. Rent/Sales are reported with different "
                "lags and are not always on the same quarter -- merged independently "
                "by their own reporting quarter."
            ),
            "indicators": indicators,
        }

        out_dir = CACHE_DIR / "housing" / "regions"
        out_dir.mkdir(parents=True, exist_ok=True)
        with open(out_dir / "narrabri_sales_rent.json", "w", encoding="utf-8") as f:
            json.dump(out, f, indent=2)

        for key, entry in indicators.items():
            vals = entry["values"]
            if vals:
                latest = max(vals, key=int)
                self.log.info(f"  [Narrabri:{key}] {len(vals)} years, latest ({latest}) = {vals[latest]}")
        self.result.towns_ok.append("Narrabri:sales_rent")


if __name__ == "__main__":
    from logger import get_logger
    get_logger()
    result = NarrabriSalesRentFetcher().run()
    sys.exit(0 if result.success else 1)
