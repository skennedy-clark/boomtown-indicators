"""
regional-indicators/fetchers/fetch_schools.py

Fetches school enrolments and teaching staff for each town from ACARA's
School Profile data, the source for the Exogenous sheet's Education
section ("Full Time Equivalent Enrolments" and "Full Time Equivalent
Teaching Staff" under each town).

Source:
    ACARA (Australian Curriculum, Assessment and Reporting Authority)
    publishes School Profile workbooks through its Data Access Program:
      - "School Profile 2008-<year>.xlsx"  every school, every year
        (about 30 MB, ~170,000 rows in the 2008-2025 file)
      - "School Profile <year>.xlsx"       the latest year only (about
        2 MB)
    Both have a "DataDictionary" sheet and one data sheet whose name
    starts "SchoolProfile". One row per school per year. The dataset is
    national and covers every state.

    Fields used (exact header text):
      "Calendar Year", "Postcode", "School Name",
      "Full Time Equivalent Enrolments",
      "Full Time Equivalent Teaching Staff"

Method:
    A town's figure for a year is the sum over every school whose
    Postcode equals the town's primary postcode (towns.toml `postcode`,
    singular).
      - Postcode is text in ACARA's file ("4405") and is compared as a
        string.
      - Only the primary postcode is used, not `postcodes` (plural).
        Toowoomba is the one town with two: 4350 alone gives 23,463.7
        FTE enrolments for 2025, which equals the reference workbook's
        figure; 4350 and 4352 together give 27,968.5. The Income sheet
        uses the same rule.
      - A school with no figure for a measure adds nothing.
      - Sums are rounded to 1 decimal place (ACARA's own precision).

Validation:
      - 2025: for all 12 towns on the sheet, both measures (24 figures)
        equal the reference workbook exactly.
      - 2024: the 2008-2025 file reproduces the sheet's existing 2024
        column exactly for all 12 towns, both measures. Reproducing the
        previous year is the cross-validation required before a new
        year is accepted.
      - Earlier years differ slightly from the sheet for some towns
        (Chinchilla, Goondiwindi, Wandoan, Toowoomba; mostly under 2%,
        up to 15% for Wandoan's small numbers), which indicates that
        ACARA has restated its back series since those columns were
        filled. The writer reports these differences and does not
        change the sheet.

Finding the file:
    Nothing year-specific is hardcoded. File names are predictable, so
    the newest is found by trying the current year, then the two years
    before it:
        School Profile 2008-<year>.xlsx      (preferred: carries history)
        School Profile <year>.xlsx           (fallback: latest year only)
    The first that exists is downloaded into cache/schools/ and reused
    on later runs (use --force to download again).

If the download fails:
    Download "School Profile" manually from ACARA's Data Access page
    (https://www.acara.edu.au/contact-us/acara-data-access), save it
    into regional-indicators/cache/schools/ keeping "School Profile"
    and the year in the file name, and re-run. The newest such file in
    that folder is used whenever the live download is unavailable.

Output: cache/schools/<slug>_schools.json, one per town with a postcode
    {"town": "Dalby", "state": "QLD", "postcode": "4405",
     "source": "...", "file": "...", "year": 2025,
     "schools": ["Dalby State School", ...],          # latest year
     "indicators": {
       "fte_enrolments":     {"2008": 2651.6, ..., "2025": 3181.4},
       "fte_teaching_staff": {"2008": 198.7,  ..., "2025": 236.0}}}
    Read by transform/xlsx_update/update_education.py.

Notes:
    Reading the 30 MB file takes roughly half a minute.

Usage:
    python run_update.py --only schools
    python fetchers/fetch_schools.py            # standalone
"""

from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).parent.parent))

from fetchers.base import BaseFetcher
from config import CACHE_DIR

try:
    import requests
    import openpyxl
except ImportError:
    raise ImportError("pip install requests openpyxl")


BASE_URL = "https://dataandreporting.blob.core.windows.net/anrdataportal/Data-Access-Program/"
ACARA_PAGE = "https://www.acara.edu.au/contact-us/acara-data-access"
HISTORY_START_YEAR = 2008          # first year of ACARA's multi-year file
YEARS_BACK_TO_TRY = 3
TIMEOUT_S = 300

SCHOOLS_CACHE_DIR = CACHE_DIR / "schools"

COL_YEAR      = "Calendar Year"
COL_POSTCODE  = "Postcode"
COL_NAME      = "School Name"
COL_ENROL     = "Full Time Equivalent Enrolments"
COL_STAFF     = "Full Time Equivalent Teaching Staff"
REQUIRED_COLUMNS = (COL_YEAR, COL_POSTCODE, COL_NAME, COL_ENROL, COL_STAFF)

LOCAL_NAME_RE = re.compile(r"school\s*profile.*?(\d{4})(?!.*\d{4})", re.IGNORECASE)


def candidate_file_names(this_year: int) -> list[tuple[str, int]]:
    """Return (file name, latest year in it) candidates, newest first, with
    the multi-year file before the single-year one for each year.
    """
    names = []
    for year in range(this_year, this_year - YEARS_BACK_TO_TRY, -1):
        names.append((f"School Profile {HISTORY_START_YEAR}-{year}.xlsx", year))
        names.append((f"School Profile {year}.xlsx", year))
    return names


def newest_local_file() -> tuple[Path, int] | None:
    """Return (path, year) for the newest School Profile workbook in
    cache/schools/, or None.

    Files are ranked by the last year in their name. For the same year
    the larger file is preferred, which selects the multi-year file over
    the single-year one.
    """
    best: tuple[Path, int] | None = None
    if SCHOOLS_CACHE_DIR.exists():
        for path in SCHOOLS_CACHE_DIR.glob("*.xlsx"):
            m = LOCAL_NAME_RE.search(path.stem)
            if not m:
                continue
            year = int(m.group(1))
            if (best is None or year > best[1]
                    or (year == best[1] and path.stat().st_size > best[0].stat().st_size)):
                best = (path, year)
    return best


def _number(value) -> float:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else 0.0


def aggregate_by_postcode(xlsx_path: Path, postcodes: set[str]) -> dict:
    """Read a School Profile workbook and return
    {postcode: {year: {"enrol": sum, "staff": sum, "schools": [names]}}}
    for the requested postcodes.

    The sheet is streamed (read_only) because the multi-year file is
    too large to load whole.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = next(
            (ws for ws in wb.worksheets if ws.title.lower().startswith("schoolprofile")), None
        )
        if sheet is None:
            raise ValueError(
                f"{xlsx_path.name} has no sheet starting 'SchoolProfile' "
                f"(sheets: {wb.sheetnames})"
            )

        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        missing = [c for c in REQUIRED_COLUMNS if c not in header]
        if missing:
            raise ValueError(
                f"{xlsx_path.name} is missing expected column(s) {missing} -- "
                f"ACARA may have renamed them."
            )
        i_year, i_pc, i_name, i_enrol, i_staff = (header.index(c) for c in REQUIRED_COLUMNS)

        result: dict = {}
        for row in rows:
            postcode = str(row[i_pc]).strip() if row[i_pc] is not None else ""
            if postcode not in postcodes:
                continue
            year = row[i_year]
            if not isinstance(year, (int, float)):
                continue
            cell = result.setdefault(postcode, {}).setdefault(
                int(year), {"enrol": 0.0, "staff": 0.0, "schools": []}
            )
            cell["enrol"] += _number(row[i_enrol])
            cell["staff"] += _number(row[i_staff])
            cell["schools"].append(str(row[i_name]).strip())
        return result
    finally:
        wb.close()


class ACARASchoolsFetcher(BaseFetcher):

    SOURCE_NAME      = "schools"
    SUPPORTED_STATES = []   # national dataset

    def fetch_all(self):
        towns = [t for t in self.applicable_towns() if t.postcode]
        if not towns:
            self.log.info("No towns with a postcode configured — nothing to fetch")
            return

        SCHOOLS_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        xlsx_path = self._get_workbook()
        if xlsx_path is None:
            self.result.add_error(
                "ALL",
                "Could not download ACARA's School Profile and no copy is saved locally. "
                f"Download it from {ACARA_PAGE} in a browser, save it into "
                f"{SCHOOLS_CACHE_DIR} keeping 'School Profile' and the year in the file "
                "name, then re-run."
            )
            return

        self.log.info(f"  Reading {xlsx_path.name} (the multi-year file takes about half a minute)...")
        try:
            data = aggregate_by_postcode(xlsx_path, {str(t.postcode) for t in towns})
        except ValueError as exc:
            self.result.add_error("ALL", str(exc))
            return

        for town in towns:
            by_year = data.get(str(town.postcode))
            if not by_year:
                self.log.warning(f"  [{town.name}] no schools found for postcode {town.postcode}")
                self.result.towns_skipped.append(town.name)
                continue

            latest = max(by_year)
            out = {
                "town":     town.name,
                "state":    town.state,
                "postcode": str(town.postcode),
                "source":   "ACARA School Profile (Data Access Program)",
                "file":     xlsx_path.name,
                "method":   "sum over all schools whose Postcode equals the town's primary postcode",
                "year":     latest,
                "schools":  sorted(by_year[latest]["schools"]),
                "indicators": {
                    "fte_enrolments":     {str(y): round(by_year[y]["enrol"], 1) for y in sorted(by_year)},
                    "fte_teaching_staff": {str(y): round(by_year[y]["staff"], 1) for y in sorted(by_year)},
                },
            }
            out_path = SCHOOLS_CACHE_DIR / f"{town.slug}_schools.json"
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(out, f, indent=2)

            self.log.info(
                f"  {town.name} (postcode {town.postcode}): {len(out['schools'])} schools, "
                f"{latest} FTE enrolments = {out['indicators']['fte_enrolments'][str(latest)]:,}, "
                f"FTE teaching staff = {out['indicators']['fte_teaching_staff'][str(latest)]:,} "
                f"({min(by_year)}-{latest})"
            )
            self.result.towns_ok.append(town.name)

    def _get_workbook(self) -> Path | None:
        """Return the path of the newest School Profile workbook, or None.

        Candidates are tried newest first: a cached copy is used if
        present (unless --force), otherwise the file is downloaded if
        ACARA has it. If none can be obtained, the newest workbook
        already saved locally is used.
        """
        for name, year in candidate_file_names(datetime.now().year):
            url = BASE_URL + quote(name)
            dest = SCHOOLS_CACHE_DIR / name
            try:
                if dest.exists() and not self.force:
                    self.log.info(f"  Using cached {name}")
                    return dest
                head = requests.head(url, timeout=60)
                if head.status_code != 200:
                    continue
                self.log.info(f"  Downloading {name} ...")
                resp = requests.get(url, timeout=TIMEOUT_S)
                resp.raise_for_status()
                if not resp.content.startswith(b"PK"):
                    raise ValueError("response is not an xlsx file")
                dest.write_bytes(resp.content)
                self.log.info(f"  Saved {name} ({len(resp.content) / 1_048_576:.1f} MB)")
                return dest
            except (requests.RequestException, ValueError) as exc:
                self.log.warning(f"  Could not get {name}: {exc}")

        local = newest_local_file()
        if local:
            self.log.warning(
                f"  Live download unavailable -- using the local copy {local[0].name}. "
                f"Check ACARA hasn't published a newer year."
            )
            return local[0]
        return None


if __name__ == "__main__":
    from logger import get_logger
    get_logger()
    result = ACARASchoolsFetcher().run()
    sys.exit(0 if result.success else 1)
