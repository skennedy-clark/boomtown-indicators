"""
regional-indicators/transform/chart_audit.py

Audits the data ranges of every chart series in the indicators workbook.

Each chart series plots one row of a data sheet (Population, Employment,
Housing, Income, Business, Crime, Exogenous) across a range of year
columns. After the annual update a series is out of date when the row
holds a figure for a year that the chart does not show. This script
lists every series, the years it shows and the years its row has data
for, and classifies it. It changes nothing.

The workbook is read directly (chart definitions from the .xlsx package,
cell values with openpyxl), so Excel is not required.

Chart filters: most charts in the workbook use Excel's chart filters.
A series then has a full range (for example Housing!C51:AE51) and shows
only the columns whose categories are not filtered out (for example
C51:AA51). The full range is stored in the chart as c15:fullRef and is
what Excel reports as the series formula; the shown columns are stored
as the ordinary reference (c:f), which becomes a list of separate areas
when a column in the middle is filtered out. A year can therefore be
missing from a chart because it lies beyond the full range, or because
it lies inside the full range but is filtered out.

Status of a series:
    OK              the chart shows every year up to the last year of
                    the row's data
    EXTEND          the row has data beyond the end of the full range, or
                    the series' category (axis label) range stops short of
                    the data; a longer range is proposed (and any
                    filtered-out years with data are listed to be shown)
    UNHIDE          the full range already covers the new data, but the
                    years are filtered out; the years to show are listed
    PROJECTION      the chart already shows figures for years after the
                    cut-off year; listed for review, no change proposed
    ENDS_PAST_DATA  no year with data is missing, but the last year
                    shown has no figure
    MULTI_AREA      the full range itself is made of several separate
                    areas; no change is proposed
    UNRESOLVED      the range could not be interpreted (not a single
                    row, no year header found above it, or a sheet that
                    is not in the workbook)

How years are found: the year header for a data row is the nearest row
above it that holds a run of at least five consecutive years, either as
calendar years (2025) or as financial-year labels ("2024/25", read as
2025). Sections with their own header row (for example Fuel and
Education on the Exogenous sheet) are therefore read against that header
and not against the top of the sheet.

Spacer columns: a shown column after the row's figures that has neither
a year in the header nor a figure draws an empty, unlabelled slot at the
end of the chart's axis. Its position in the category range is listed
(blank_positions) so that the chart step can filter it out when no
series in the chart has a figure there.

Which years are missing: starting from the last shown column that has a
figure, the run of numeric cells is followed to the right until the
first empty cell or the cut-off year. The years in that run that the
chart does not show are missing. Data that resumes after a gap is
mentioned in the note and is not included. When the full range is
extended, the category (axis label) range is extended with it if it
ends on the same column.

Fixed value-axis limits: a chart whose value axis has a fixed maximum
(or minimum) that its data goes beyond draws the bars or lines cut off
at the edge of the plot. Such charts are listed separately
(axis_limits); for a stacked chart the height of each stack is compared.
The chart step (update_charts.py) sets those limits back to automatic.
A fixed limit can come from a template page (transform/add_town.py) or
be outgrown by a new year's figure.

With --reference <workbook>, the same audit is run on a second workbook
and each series is listed with the columns it shows there. Series are paired
by page, chart position on the page and series order.

Output: a CSV file (default chart_audit.csv), one line per series, and
a summary on the console.

Usage:
    python chart_audit.py <workbook.xlsx> [--out chart_audit.csv]
                          [--last-year 2025]
                          [--reference <other workbook.xlsx>]
"""

from __future__ import annotations

import argparse
import csv
import posixpath
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

NS = {
    "c":   "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a":   "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "r":   "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m":   "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "c15": "http://schemas.microsoft.com/office/drawing/2012/chart",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
CHART_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
DRAWING_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"

MIN_YEARS_IN_HEADER = 5
HEADER_SEARCH_ROWS = 400

STATUSES = ("OK", "EXTEND", "UNHIDE", "PROJECTION", "ENDS_PAST_DATA", "MULTI_AREA", "UNRESOLVED")
CHANGE_STATUSES = ("EXTEND", "UNHIDE")

CSV_COLUMNS = [
    "page", "chart_at", "chart_title", "series_no", "series_name",
    "data_sheet", "row", "row_label", "sub_label", "block",
    "values", "categories", "shown_values", "range_from_year", "range_to_year",
    "data_to_year", "status", "proposed_values", "proposed_categories",
    "unhide_years", "blank_positions", "reference_values", "note",
]


# ── Workbook package ───────────────────────────────────────────────────────────

@dataclass
class Area:
    """One rectangular area of a range reference."""
    sheet: str
    col1: int
    row1: int
    col2: int
    row2: int

    @property
    def single_row(self) -> bool:
        return self.row1 == self.row2

    def with_end_column(self, col: int) -> str:
        return format_area(self.sheet, self.col1, self.row1, col, self.row2)


@dataclass
class Series:
    page: str
    chart_at: str               # top-left cell of the chart on its page
    chart_order: tuple          # (row, column) of that cell, for sorting
    chart_title: str
    series_no: int
    series_name: str
    values: str                 # full range (what Excel reports in the series formula)
    categories: str             # full category range
    shown_values: str = ""      # the columns the chart shows (differs when a chart filter is set)
    shown_categories: str = ""
    # filled in by assess()
    data_sheet: str = ""
    row: int | None = None
    row_label: str = ""
    sub_label: str = ""
    block: str = ""
    range_from_year: int | None = None
    range_to_year: int | None = None
    data_to_year: int | None = None
    status: str = ""
    proposed_values: str = ""
    proposed_categories: str = ""
    unhide_years: list[int] = field(default_factory=list)
    blank_positions: list[int] = field(default_factory=list)   # shown categories with no year and no figure
    reference_values: str = ""
    notes: list[str] = field(default_factory=list)

    def as_row(self) -> dict:
        row = {name: getattr(self, name, "") for name in CSV_COLUMNS if name != "note"}
        row["note"] = "; ".join(self.notes)
        row["unhide_years"] = " ".join(str(y) for y in self.unhide_years)
        row["blank_positions"] = " ".join(str(n) for n in self.blank_positions)
        return {k: ("" if v is None else v) for k, v in row.items()}


def _relationships(package: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """Return {relationship id: (type, target part)} for a package part."""
    folder, name = posixpath.split(part)
    rels_part = posixpath.join(folder, "_rels", name + ".rels")
    if rels_part not in package.namelist():
        return {}
    out = {}
    for rel in ET.fromstring(package.read(rels_part)).findall("rel:Relationship", NS):
        target = rel.get("Target")
        if rel.get("TargetMode") == "External":
            continue
        resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        out[rel.get("Id")] = (rel.get("Type"), resolved)
    return out


def _text(element, path: str) -> str:
    if element is None:
        return ""
    return "".join(node.text or "" for node in element.findall(path, NS)).strip()


def _reference(series_element, *paths: str) -> str:
    for path in paths:
        node = series_element.find(path, NS)
        if node is not None and node.text:
            return node.text.strip()
    return ""


VALUE_PATHS = ("c:val/c:numRef", "c:yVal/c:numRef")
CATEGORY_PATHS = ("c:cat/c:numRef", "c:cat/c:strRef", "c:cat/c:multiLvlStrRef",
                  "c:xVal/c:numRef", "c:xVal/c:strRef")


def _ranges(series_element, paths) -> tuple[str, str]:
    """Return (shown reference, full reference) for the values or the
    categories of a series. Without a chart filter the two are equal."""
    for path in paths:
        node = series_element.find(path, NS)
        if node is None:
            continue
        shown = (node.findtext("c:f", "", NS) or "").strip()
        full = (node.findtext("c:extLst/c:ext/c15:fullRef/c15:sqref", "", NS) or "").strip()
        if shown or full:
            return shown or full, full or shown
    return "", ""


def read_chart_series(workbook_path: Path) -> list[Series]:
    """Return every chart series in the workbook, in page and chart order."""
    series_list: list[Series] = []
    with zipfile.ZipFile(workbook_path) as package:
        workbook_rels = _relationships(package, "xl/workbook.xml")
        workbook = ET.fromstring(package.read("xl/workbook.xml"))
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            page = sheet.get("name")
            sheet_part = workbook_rels[sheet.get(f"{{{NS['r']}}}id")][1]
            for rel_type, drawing_part in _relationships(package, sheet_part).values():
                if rel_type != DRAWING_REL:
                    continue
                drawing_rels = _relationships(package, drawing_part)
                drawing = ET.fromstring(package.read(drawing_part))
                for anchor in list(drawing):
                    chart_ref = anchor.find(".//c:chart", NS)
                    origin = anchor.find("xdr:from", NS)
                    if chart_ref is None:
                        continue
                    rel_id = chart_ref.get(f"{{{NS['r']}}}id")
                    if rel_id not in drawing_rels or drawing_rels[rel_id][0] != CHART_REL:
                        continue
                    if origin is not None:
                        col = int(origin.findtext("xdr:col", "0", NS)) + 1
                        row = int(origin.findtext("xdr:row", "0", NS)) + 1
                    else:
                        col, row = 1, 1
                    chart = ET.fromstring(package.read(drawing_rels[rel_id][1]))
                    title = _text(chart.find("c:chart/c:title", NS), ".//a:t")
                    for number, ser in enumerate(chart.findall(".//c:ser", NS), start=1):
                        name = _text(ser.find("c:tx", NS), ".//c:v")
                        shown_values, full_values = _ranges(ser, VALUE_PATHS)
                        shown_categories, full_categories = _ranges(ser, CATEGORY_PATHS)
                        series_list.append(Series(
                            page=page,
                            chart_at=f"{get_column_letter(col)}{row}",
                            chart_order=(row, col),
                            chart_title=title,
                            series_no=number,
                            series_name=name,
                            values=full_values,
                            categories=full_categories,
                            shown_values=shown_values,
                            shown_categories=shown_categories,
                        ))
    return series_list


@dataclass
class AxisLimit:
    """A chart whose fixed value-axis limit cuts off its data."""
    page: str
    chart_at: str
    chart_title: str
    axis_group: int                 # 1 primary, 2 secondary (Excel's XlAxisGroup)
    minimum: float | None           # the fixed limits, None where automatic
    maximum: float | None
    data_low: float
    data_peak: float

    @property
    def label(self) -> str:
        return f"{self.page} {self.chart_at}"

    def describe(self) -> str:
        parts = []
        if self.maximum is not None and self.data_peak > self.maximum:
            parts.append(f"data reach {self.data_peak:,.0f}, above the fixed maximum {self.maximum:,.0f}")
        if self.minimum is not None and self.data_low < self.minimum:
            parts.append(f"data fall to {self.data_low:,.0f}, below the fixed minimum {self.minimum:,.0f}")
        axis = "" if self.axis_group == 1 else " (secondary axis)"
        return f"{self.label} '{self.chart_title}'{axis}: " + "; ".join(parts)


AXIS_GROUPS = ("barChart", "bar3DChart", "lineChart", "line3DChart", "areaChart", "area3DChart",
               "scatterChart", "radarChart")


def _number(text: str | None) -> float | None:
    try:
        return float(text) if text not in (None, "") else None
    except ValueError:
        return None


def axis_limits(workbook_path: Path, pages: set[str] | None = None) -> list[AxisLimit]:
    """Every chart (optionally on `pages`) whose fixed value-axis maximum
    or minimum its data go beyond. The data are the columns each series
    shows; for a stacked group, the positive figures of each category are
    added up."""
    found: list[AxisLimit] = []
    pending = []                                   # (page, chart_at, title, group, min, max, [references], stacked)
    sheets: set[str] = set()
    with zipfile.ZipFile(workbook_path) as package:
        workbook_rels = _relationships(package, "xl/workbook.xml")
        workbook = ET.fromstring(package.read("xl/workbook.xml"))
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            page = sheet.get("name")
            if pages is not None and page not in pages:
                continue
            sheet_part = workbook_rels[sheet.get(f"{{{NS['r']}}}id")][1]
            for rel_type, drawing_part in _relationships(package, sheet_part).values():
                if rel_type != DRAWING_REL:
                    continue
                drawing_rels = _relationships(package, drawing_part)
                for anchor in list(ET.fromstring(package.read(drawing_part))):
                    chart_ref = anchor.find(".//c:chart", NS)
                    origin = anchor.find("xdr:from", NS)
                    if chart_ref is None:
                        continue
                    rel_id = chart_ref.get(f"{{{NS['r']}}}id")
                    if rel_id not in drawing_rels or drawing_rels[rel_id][0] != CHART_REL:
                        continue
                    col = int(origin.findtext("xdr:col", "0", NS)) + 1 if origin is not None else 1
                    row = int(origin.findtext("xdr:row", "0", NS)) + 1 if origin is not None else 1
                    chart = ET.fromstring(package.read(drawing_rels[rel_id][1]))
                    plot = chart.find("c:chart/c:plotArea", NS)
                    if plot is None:
                        continue
                    title = _text(chart.find("c:chart/c:title", NS), ".//a:t")
                    value_axes = {}
                    for order, axis in enumerate(plot.findall("c:valAx", NS)):
                        scaling = axis.find("c:scaling", NS)
                        value_axes[axis.find("c:axId", NS).get("val")] = (
                            order,
                            _number(scaling.find("c:min", NS).get("val")) if scaling.find("c:min", NS) is not None else None,
                            _number(scaling.find("c:max", NS).get("val")) if scaling.find("c:max", NS) is not None else None,
                        )
                    primary = min((o for o, _, _ in value_axes.values()), default=0)
                    for group in plot:
                        if group.tag.split("}")[1] not in AXIS_GROUPS:
                            continue
                        axis_ids = [a.get("val") for a in group.findall("c:axId", NS)]
                        axis = next((value_axes[i] for i in axis_ids if i in value_axes), None)
                        if axis is None or (axis[1] is None and axis[2] is None):
                            continue
                        grouping = group.find("c:grouping", NS)
                        stacked = grouping is not None and "stacked" in (grouping.get("val") or "").lower()
                        references = [_ranges(ser, VALUE_PATHS)[0] for ser in group.findall("c:ser", NS)]
                        for reference in references:
                            for area in parse_reference(reference) or []:
                                sheets.add(area.sheet)
                        pending.append((page, f"{get_column_letter(col)}{row}", title,
                                        1 if axis[0] == primary else 2, axis[1], axis[2], references, stacked))
    if not pending:
        return found
    cells = Cells(workbook_path, sheets)
    for page, chart_at, title, axis_group, minimum, maximum, references, stacked in pending:
        columns: list[list[float]] = []
        for reference in references:
            figures = []
            for area in parse_reference(reference) or []:
                if not cells.has(area.sheet):
                    continue
                for r in range(area.row1, area.row2 + 1):
                    for c in range(area.col1, area.col2 + 1):
                        value = cells.value(area.sheet, r, c)
                        figures.append(float(value) if is_number(value) else None)
            columns.append(figures)
        numbers = [v for figures in columns for v in figures if v is not None]
        if not numbers:
            continue
        if stacked:
            width = max(len(f) for f in columns)
            peak = max(sum(f[i] for f in columns if i < len(f) and f[i] is not None and f[i] > 0)
                       for i in range(width))
        else:
            peak = max(numbers)
        low = min(numbers)
        if (maximum is not None and peak > maximum) or (minimum is not None and low < minimum):
            found.append(AxisLimit(page, chart_at, title, axis_group, minimum, maximum, low, peak))
    found.sort(key=lambda a: (a.page, a.chart_at))
    return found


# ── Range references ───────────────────────────────────────────────────────────

_AREA = re.compile(
    r"^(?:'((?:[^']|'')+)'|([^'!,()]+))!"
    r"\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?$"
)


def parse_reference(reference: str) -> list[Area] | None:
    """Split a chart reference into its areas, or return None if any
    part is not a plain sheet range (for example a defined name)."""
    text = reference.strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    if not text:
        return None
    areas = []
    for part in _split_areas(text):
        match = _AREA.match(part.strip())
        if not match:
            return None
        quoted, plain, col1, row1, col2, row2 = match.groups()
        sheet = quoted.replace("''", "'") if quoted is not None else plain
        c1 = column_index_from_string(col1.upper())
        c2 = column_index_from_string(col2.upper()) if col2 else c1
        r1 = int(row1)
        r2 = int(row2) if row2 else r1
        areas.append(Area(sheet, min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)))
    return areas


def _split_areas(text: str) -> list[str]:
    """Split on commas that are outside quoted sheet names."""
    parts, current, quoted = [], [], False
    for char in text:
        if char == "'":
            quoted = not quoted
        if char == "," and not quoted:
            parts.append("".join(current))
            current = []
        else:
            current.append(char)
    parts.append("".join(current))
    return parts


def format_area(sheet: str, col1: int, row1: int, col2: int, row2: int) -> str:
    name = sheet if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", sheet) else "'" + sheet.replace("'", "''") + "'"
    start = f"${get_column_letter(col1)}${row1}"
    end = f"${get_column_letter(col2)}${row2}"
    return f"{name}!{start}" if start == end else f"{name}!{start}:{end}"


# ── Cell values ────────────────────────────────────────────────────────────────

class Cells:
    """Cached values of the data sheets, read once with openpyxl."""

    def __init__(self, workbook_path: Path, sheets: set[str]):
        book = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
        self.sheets: dict[str, list[tuple]] = {}
        for name in sheets:
            if name in book.sheetnames:
                self.sheets[name] = [row for row in book[name].iter_rows(values_only=True)]
        book.close()

    def has(self, sheet: str) -> bool:
        return sheet in self.sheets

    def value(self, sheet: str, row: int, col: int):
        rows = self.sheets[sheet]
        if row < 1 or row > len(rows):
            return None
        cells = rows[row - 1]
        return cells[col - 1] if 1 <= col <= len(cells) else None

    def width(self, sheet: str, row: int) -> int:
        rows = self.sheets[sheet]
        return len(rows[row - 1]) if 1 <= row <= len(rows) else 0


def is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


_FINANCIAL_YEAR = re.compile(r"^\s*((?:19|20)\d{2})\s*[/-]\s*(\d{2}|\d{4})\s*$")


def as_year(value) -> int | None:
    """Return the year a header cell stands for, or None.

    Accepts a calendar year (2025) and a financial-year label
    ("2024/25" or "2024-25"), which counts as the year it ends in.
    """
    if is_number(value) and float(value).is_integer() and 1990 <= value <= 2100:
        return int(value)
    if isinstance(value, str):
        match = _FINANCIAL_YEAR.match(value)
        if match:
            start = int(match.group(1))
            end = int(match.group(2))
            if end == (start + 1) % 100 or end == start + 1:
                return start + 1
    return None


def is_year_row(cells: "Cells", sheet: str, row: int, columns) -> bool:
    """True if the row holds a run of at least MIN_YEARS_IN_HEADER
    consecutive years in adjacent cells of `columns`. Requiring a run
    excludes data rows whose values happen to look like years.
    """
    run, previous = 0, None
    for col in columns:
        year = as_year(cells.value(sheet, row, col))
        if year is not None and previous is not None and year == previous + 1:
            run += 1
        else:
            run = 1 if year is not None else 0
        if run >= MIN_YEARS_IN_HEADER:
            return True
        previous = year
    return False


def find_year_header(cells: Cells, sheet: str, row: int, col1: int, col2: int) -> int | None:
    """Return the nearest row above `row` that is a year header.

    A row qualifies if it holds a run of consecutive years (see
    is_year_row), looked for in the range's own columns when the range
    is wide enough and across the whole row otherwise. Where a sheet has
    a row of financial-year labels above a row of calendar years, the
    lower row is the one returned.
    """
    span = col2 - col1 + 1
    for candidate in range(row - 1, max(0, row - 1 - HEADER_SEARCH_ROWS), -1):
        if span >= MIN_YEARS_IN_HEADER:
            columns = range(col1, col2 + 1)
        else:
            columns = range(1, cells.width(sheet, candidate) + 1)
        if is_year_row(cells, sheet, candidate, columns):
            return candidate
    return None


def find_block(cells: Cells, sheet: str, row: int) -> str:
    """Return the heading of the block a data row belongs to: the nearest
    row above with text in column A and no numbers in the row."""
    for candidate in range(row - 1, max(0, row - 60), -1):
        label = cells.value(sheet, candidate, 1)
        if not isinstance(label, str) or not label.strip():
            continue
        width = cells.width(sheet, candidate)
        if not any(is_number(cells.value(sheet, candidate, col)) for col in range(2, width + 1)):
            return label.strip()
    return ""


# ── Assessment ─────────────────────────────────────────────────────────────────

def assess(series: Series, cells: Cells, last_year: int) -> None:
    """Classify one series and fill in its descriptive fields."""
    areas = parse_reference(series.values)
    if not areas:
        series.status = "UNRESOLVED"
        series.notes.append("values reference is not a sheet range")
        return
    area = areas[0]
    series.data_sheet = area.sheet
    if not cells.has(area.sheet):
        series.status = "UNRESOLVED"
        series.notes.append(f"sheet '{area.sheet}' is not in the workbook")
        return
    if len(areas) > 1:
        series.status = "MULTI_AREA"
        series.notes.append(f"the full range has {len(areas)} separate areas")
        return
    if not area.single_row:
        series.status = "UNRESOLVED"
        series.notes.append("values range is not a single row")
        return

    row = area.row1
    series.row = row
    label = cells.value(area.sheet, row, 1)
    sub_label = cells.value(area.sheet, row, 2)
    series.row_label = str(label).strip() if label is not None else ""
    series.sub_label = str(sub_label).strip() if isinstance(sub_label, str) else ""
    series.block = find_block(cells, area.sheet, row)

    header = find_year_header(cells, area.sheet, row, area.col1, area.col2)
    if header is None:
        series.status = "UNRESOLVED"
        series.notes.append("no year header found above the row")
        return

    def year_at(col: int) -> int | None:
        return as_year(cells.value(area.sheet, header, col))

    def has_data(col: int) -> bool:
        return is_number(cells.value(area.sheet, row, col))

    shown = _shown_columns(series.shown_values, area)
    if series.shown_values and normalise_reference(series.shown_values) != normalise_reference(series.values):
        series.notes.append("chart filter: " + describe_columns(sorted(shown), area.col1, area.col2))
    shown_end = max(shown)

    # Shown columns after the row's figures that have neither a year in
    # the header nor a figure (spacer columns). They draw an empty, unlabelled
    # slot at the end of the axis; their positions in the category range
    # are recorded so that the chart can filter them out.
    # The label is read from the chart's own category cells where they are a
    # single row (the Income sheet's charts label their axis from the fiscal
    # years in row 1, while its data rows sit under calendar years in row 2).
    # A row with no figures at all (a town just added, not yet written) has
    # no spacer columns to report.
    cat_areas = parse_reference(series.categories) if series.categories else None
    single_cat = cat_areas[0] if cat_areas and len(cat_areas) == 1 and cat_areas[0].row1 == cat_areas[0].row2 else None
    cat_start = single_cat.col1 if single_cat else area.col1

    def labelled(col: int) -> bool:
        if year_at(col) is not None:
            return True
        if single_cat and cells.has(single_cat.sheet):
            value = cells.value(single_cat.sheet, single_cat.row1, col)
            return value not in (None, "") and not (isinstance(value, str) and not value.strip())
        return False

    last_figure = max((c for c in shown if has_data(c)), default=0)
    series.blank_positions = [] if not last_figure else [
        c - cat_start + 1 for c in sorted(shown)
        if c > last_figure and not labelled(c) and not has_data(c) and c >= cat_start
    ]
    shown_years = [year_at(c) for c in sorted(shown) if year_at(c) is not None]
    series.range_from_year = shown_years[0] if shown_years else None
    series.range_to_year = shown_years[-1] if shown_years else None

    shown_with_data = [c for c in sorted(shown) if has_data(c)]
    if shown_with_data and (year_at(shown_with_data[-1]) or 0) > last_year:
        series.status = "PROJECTION"
        series.data_to_year = year_at(shown_with_data[-1])
        series.notes.append(f"the chart already shows figures for years after the cut-off year {last_year}")
        return

    # The run of data that continues from the last shown column with a figure.
    start = _last_data_column(has_data, area.col1, shown_end, shown)
    if start is None:
        series.status = "ENDS_PAST_DATA"
        series.notes.append("the row has no figures in the years shown")
        return
    width = max(cells.width(area.sheet, row), cells.width(area.sheet, header))
    run_end = start
    while run_end + 1 <= width and has_data(run_end + 1) and year_at(run_end + 1) is not None \
            and year_at(run_end + 1) <= last_year:
        run_end += 1
    series.data_to_year = year_at(run_end)
    if run_end + 1 <= width and has_data(run_end + 1) and (year_at(run_end + 1) or 0) > last_year:
        series.notes.append(f"the row continues after the cut-off year {last_year} (projections), not included")
    later = [c for c in range(run_end + 2, width + 1) if has_data(c) and (year_at(c) or 0) and year_at(c) <= last_year]
    if later:
        series.notes.append(
            f"more data after a gap, from {year_at(later[0])} "
            f"({get_column_letter(later[0])}{row}), not included"
        )

    missing = [c for c in range(start + 1, run_end + 1) if c not in shown]
    if not missing:
        if not has_data(shown_end):
            series.status = "ENDS_PAST_DATA"
            series.notes.append(
                f"the last year shown, {series.range_to_year}, has no figure; "
                f"the row's last figure is for {series.data_to_year}"
            )
        else:
            series.status = "OK"
        return

    series.unhide_years = [year_at(c) for c in missing if c <= area.col2]
    if run_end > area.col2:
        series.status = "EXTEND"
        series.proposed_values = area.with_end_column(run_end)
        series.unhide_years = [year_at(c) for c in missing]
        note = _check_categories(series, area, header, run_end, cells)
        if note:
            series.notes.append(note)
    else:
        series.status = "UNHIDE"
        # The years can only be shown if the series' category range
        # reaches them: in a chart with two chart groups (columns and
        # lines) each group takes its axis labels from its own series.
        cat_areas = parse_reference(series.categories) if series.categories else None
        if cat_areas and len(cat_areas) == 1 and cat_areas[0].single_row and cat_areas[0].col2 < run_end:
            cat = cat_areas[0]
            series.status = "EXTEND"
            series.proposed_values = series.values
            series.proposed_categories = cat.with_end_column(area.col2)
            series.notes.append(
                f"category range ends at column {get_column_letter(cat.col2)}, before the data; "
                f"extended to {get_column_letter(area.col2)} to match the values"
            )


def _shown_columns(shown_reference: str, area: Area) -> set[int]:
    """Columns of the full range that the chart shows."""
    shown_areas = parse_reference(shown_reference) if shown_reference else None
    columns = set()
    for part in shown_areas or []:
        if part.sheet == area.sheet and part.row1 == area.row1 and part.single_row:
            columns.update(range(part.col1, part.col2 + 1))
    return columns or set(range(area.col1, area.col2 + 1))


def _last_data_column(has_data, first: int, last: int, allowed: set[int] | None = None) -> int | None:
    for col in range(last, first - 1, -1):
        if (allowed is None or col in allowed) and has_data(col):
            return col
    return None


def describe_columns(columns: list[int], first: int, last: int) -> str:
    """Describe which columns of first..last are shown, e.g. 'shows C-Z, AB of C-AE'."""
    runs, start, previous = [], None, None
    for col in columns:
        if start is None:
            start = previous = col
        elif col == previous + 1:
            previous = col
        else:
            runs.append((start, previous))
            start = previous = col
    if start is not None:
        runs.append((start, previous))
    shown = ", ".join(
        get_column_letter(a) if a == b else f"{get_column_letter(a)}-{get_column_letter(b)}" for a, b in runs
    )
    return f"shows {shown} of {get_column_letter(first)}-{get_column_letter(last)}"


def normalise_reference(reference: str) -> str:
    text = (reference or "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1]
    return text.replace("$", "").replace("'", "").replace(" ", "").upper()


def _check_categories(series: Series, area: Area, header: int, run_end: int, cells: Cells) -> str:
    """Propose a category range to go with an extended values range.

    The category range is extended when it covers the same columns as
    the values. A category range that is already wider (room left for
    projections) is kept unless the values would run past it.
    """
    if not series.categories:
        return "no category range"
    cat_areas = parse_reference(series.categories)
    if not cat_areas:
        return "category reference is not a sheet range"
    if len(cat_areas) > 1:
        return "category range has several areas"
    cat = cat_areas[0]
    if cat.col1 != area.col1 or cat.col2 < area.col2:
        if cat.col2 < run_end:
            series.proposed_categories = cat.with_end_column(run_end)
        return (
            f"category range covers columns {get_column_letter(cat.col1)}-{get_column_letter(cat.col2)}, "
            f"values cover {get_column_letter(area.col1)}-{get_column_letter(area.col2)}"
        )
    if cat.col2 > area.col2:
        if run_end > cat.col2:
            series.proposed_categories = cat.with_end_column(run_end)
        return f"axis runs to column {get_column_letter(cat.col2)}, beyond the values"
    series.proposed_categories = cat.with_end_column(run_end)
    return ""


def audit(workbook_path: Path, last_year: int) -> list[Series]:
    series_list = read_chart_series(workbook_path)
    sheets = set()
    for series in series_list:
        for reference in (series.values, series.categories):
            for area in parse_reference(reference) or []:
                sheets.add(area.sheet)
    cells = Cells(workbook_path, sheets)
    for series in series_list:
        assess(series, cells, last_year)
    series_list.sort(key=lambda s: (s.page, s.chart_order, s.series_no))
    return series_list


def add_reference(series_list: list[Series], reference_list: list[Series]) -> int:
    """Record, for each series, the columns shown by the series in the
    same position in the reference workbook. Returns the number paired."""
    by_position = {(s.page, s.chart_at, s.series_no): s for s in reference_list}
    paired = 0
    for series in series_list:
        other = by_position.get((series.page, series.chart_at, series.series_no))
        if other is not None:
            series.reference_values = other.shown_values
            paired += 1
    return paired


# ── Reporting ──────────────────────────────────────────────────────────────────

def summary_lines(series_list: list[Series], last_year: int) -> list[str]:
    charts = {(s.page, s.chart_at) for s in series_list}
    pages = sorted({s.page for s in series_list})
    lines = [
        f"{len(charts)} charts, {len(series_list)} series on {len(pages)} pages. "
        f"Cut-off year: {last_year}.",
        "",
    ]
    totals = Counter(s.status for s in series_list)
    lines.append("Status          Series")
    for status in STATUSES:
        lines.append(f"  {status:<14}{totals.get(status, 0):>6}")

    lines += ["", "By data sheet:"]
    lines.append("  " + f"{'':<12}" + "".join(f"{status:>16}" for status in STATUSES))
    for sheet in sorted({s.data_sheet or "(none)" for s in series_list}):
        counts = Counter(s.status for s in series_list if (s.data_sheet or "(none)") == sheet)
        lines.append("  " + f"{sheet:<12}" + "".join(f"{counts.get(status, 0):>16}" for status in STATUSES))

    lines += ["", "By page:"]
    lines.append("  " + f"{'':<14}" + "".join(f"{status:>16}" for status in STATUSES))
    for page in pages:
        counts = Counter(s.status for s in series_list if s.page == page)
        lines.append("  " + f"{page[:13]:<14}" + "".join(f"{counts.get(status, 0):>16}" for status in STATUSES))

    changes = [s for s in series_list if s.status in CHANGE_STATUSES]
    if changes:
        moves = Counter((s.data_sheet, s.status, s.range_to_year, s.data_to_year) for s in changes)
        lines += ["", "Changes proposed (data sheet, change: last year shown -> last year of data):"]
        for (sheet, status, to_year, data_year), count in sorted(moves.items(), key=lambda kv: (kv[0][0], kv[0][1], -kv[1])):
            action = "extend range" if status == "EXTEND" else "show filtered years"
            lines.append(f"  {sheet:<12}{action:<20}{to_year} -> {data_year}   {count} series")
    return lines


def write_csv(series_list: list[Series], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        for series in series_list:
            writer.writerow(series.as_row())


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="List every chart series and whether its range needs extending.")
    parser.add_argument("workbook", type=Path, help="the indicators workbook (.xlsx)")
    parser.add_argument("--out", type=Path, default=Path("chart_audit.csv"),
                        help="CSV file to write (default: chart_audit.csv)")
    parser.add_argument("--last-year", type=int, default=datetime.now().year,
                        help="cut-off year: data after it is treated as projections "
                             "(default: the current year)")
    parser.add_argument("--reference", type=Path, default=None,
                        help="a second workbook whose ranges are shown alongside")
    args = parser.parse_args(argv)

    if not args.workbook.exists():
        print(f"Workbook not found: {args.workbook}")
        return 2

    series_list = audit(args.workbook, args.last_year)
    print(f"Workbook: {args.workbook}")
    for line in summary_lines(series_list, args.last_year):
        print(line)

    if args.reference is not None:
        if not args.reference.exists():
            print(f"Reference workbook not found: {args.reference}")
            return 2
        reference_list = audit(args.reference, args.last_year)
        paired = add_reference(series_list, reference_list)
        same = sum(1 for s in series_list
                   if s.reference_values and normalise_reference(s.reference_values) == normalise_reference(s.shown_values))
        print("")
        print(f"Reference: {args.reference}")
        print(f"  {paired} of {len(series_list)} series paired by page, chart position and series order")
        print(f"  {same} show the same columns in both workbooks")

    limits = axis_limits(args.workbook)
    print("")
    print(f"Charts whose fixed value-axis limit cuts off data: {len(limits)}")
    for limit in limits:
        print(f"  {limit.describe()}")

    write_csv(series_list, args.out)
    print("")
    print(f"Wrote {len(series_list)} lines to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
