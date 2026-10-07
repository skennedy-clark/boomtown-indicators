"""
tests/fake_xlwings_sheet.py -- a stand-in for an xlwings Sheet, backed by
openpyxl, for exercising the xlsx_update writers' row-finding and audit
logic where real Excel isn't available (Linux, CI).

READ-ONLY IN SPIRIT: writes go to an in-memory dict (`.writes`) and are
NEVER saved back to the workbook -- openpyxl must not be used to save
the real indicators workbook (see transform/xlsx_update/base.py).

Mimics only the xlwings behaviour the writers rely on:
  sheet.name
  sheet.used_range.last_cell.row / .column
  sheet.range((r1, c1), (r2, c2)).value   scalar / flat list / list of lists
  sheet.cells(r, c).value / .formula / .address / .number_format / .font
Numbers come back as float, as xlwings returns them.
"""
from __future__ import annotations

import openpyxl
from openpyxl.utils import get_column_letter


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v


class _Font:
    bold = False
    color = (0, 0, 0)


class _Cell:
    def __init__(self, sheet, r, c):
        self._s, self._r, self._c = sheet, r, c
        self.number_format = "General"
        self.font = _Font()

    @property
    def address(self):
        return f"${get_column_letter(self._c)}${self._r}"

    @property
    def formula(self):
        if (self._r, self._c) in self._s.writes:
            return str(self._s.writes[(self._r, self._c)])
        v = self._s._wf.cell(self._r, self._c).value
        return "" if v is None else str(v)

    @property
    def value(self):
        return self._s._value(self._r, self._c)

    @value.setter
    def value(self, v):
        self._s.writes[(self._r, self._c)] = v


class _Range:
    def __init__(self, sheet, a, b):
        self._s, (self._r1, self._c1), (self._r2, self._c2) = sheet, a, b

    @property
    def value(self):
        rows = [[self._s._value(r, c) for c in range(self._c1, self._c2 + 1)]
                for r in range(self._r1, self._r2 + 1)]
        if len(rows) == 1 and len(rows[0]) == 1:
            return rows[0][0]
        if len(rows) == 1:
            return rows[0]
        if len(rows[0]) == 1:
            return [r[0] for r in rows]
        return rows


class _LastCell:
    def __init__(self, ws):
        self.row, self.column = ws.max_row, ws.max_column


class _UsedRange:
    def __init__(self, ws):
        self.last_cell = _LastCell(ws)


class FakeSheet:
    def __init__(self, xlsx_path, sheet_name):
        self._wf = openpyxl.load_workbook(xlsx_path)[sheet_name]                  # formulas
        self._wv = openpyxl.load_workbook(xlsx_path, data_only=True)[sheet_name]  # cached values
        self.name = sheet_name
        self.writes: dict[tuple[int, int], object] = {}
        self.used_range = _UsedRange(self._wf)

    def _value(self, r, c):
        if (r, c) in self.writes:
            return _num(self.writes[(r, c)])
        return _num(self._wv.cell(r, c).value)

    def range(self, a, b):
        return _Range(self, a, b)

    def cells(self, r, c):
        return _Cell(self, r, c)