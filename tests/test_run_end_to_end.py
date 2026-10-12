"""
tests/test_run_end_to_end.py

Tests for the end-to-end runner (regional-indicators/run_end_to_end.py).

The runner is tested with small stand-in scripts: that it does not
overwrite the starting workbook, builds the steps in order with the
right arguments, continues when a step fails, and reports each step's
own summary. No fetching and no Excel.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "regional-indicators"))


@pytest.fixture
def log(tmp_path):
    import run_end_to_end as e2e

    log = e2e.Log(tmp_path / "run.log")
    yield log
    log.close()


def _script(tmp_path, name, body):
    path = tmp_path / name
    path.write_text(body, encoding="utf-8")
    return str(path)


def test_step_reports_the_writers_own_summary_line(tmp_path, log):
    import run_end_to_end as e2e

    script = _script(tmp_path, "ok.py", 'print("Dalby: WRITTEN 2025 = 1 -> $Z$9")\nprint("Summary: 8 written, 0 flagged for review.")\n')
    result = e2e.run_step("Fuel", [sys.executable, script], log)
    assert result["ok"] and result["summary"] == "Summary: 8 written, 0 flagged for review."


def test_failed_step_is_recorded_not_raised(tmp_path, log):
    import run_end_to_end as e2e

    script = _script(tmp_path, "boom.py", 'raise RuntimeError("no Excel here")\n')
    result = e2e.run_step("Crime", [sys.executable, script], log)
    assert not result["ok"] and "crashed" in result["summary"]
    log.close()
    assert "no Excel here" in (tmp_path / "run.log").read_text(encoding="utf-8")


def test_fetch_totals_are_combined(tmp_path, log):
    import run_end_to_end as e2e

    script = _script(tmp_path, "fetch.py", 'import sys\nprint("  ✓ Passed : 16")\nprint("  ✗ Failed : 2")\nsys.exit(1)\n')
    result = e2e.run_step("Fetch everything", [sys.executable, script], log)
    assert result["summary"] == "passed: 16, failed: 2"


def test_steps_are_fetch_then_every_writer_then_the_website(tmp_path):
    import argparse
    import run_end_to_end as e2e

    args = argparse.Namespace(skip_fetch=False, business_year=2025, last_year=2025, visible=True,
                              web_out=tmp_path / "3 Web Content", previous_web=tmp_path / "old",
                              data_out=tmp_path / "1 Data")
    steps = e2e.build_steps(args, tmp_path / "test-copy.xlsx")
    names = [name for name, _ in steps]
    assert names[:2] == ["Manual downloads", "Fetch everything"]
    assert names[-3:] == ["Charts", "Website folder", "Data folder"]
    assert len(names) == len(e2e.WRITER_STEPS) + 5

    commands = dict(steps)
    business = commands["Business"]
    assert business[-2:] == ["2025", "--visible"]                      # the only writer that takes a year
    assert Path(commands["Income"][3]).name == "ato"                   # Income reads cache/ato, not cache/income
    assert Path(commands["Population: LGA"][3]).name == "population"
    assert "--compare-with" in commands["Website folder"]

    # every writer script listed in WRITER_STEPS exists
    for _, script, _ in e2e.WRITER_STEPS:
        assert (e2e.WRITERS / script).exists(), script


def test_skip_fetch_leaves_the_fetch_step_out(tmp_path):
    import argparse
    import run_end_to_end as e2e

    args = argparse.Namespace(skip_fetch=True, business_year=2025, last_year=2025, visible=False,
                              web_out=tmp_path / "w", previous_web=None, data_out=tmp_path / "d")
    steps = e2e.build_steps(args, tmp_path / "t.xlsx")
    assert "Fetch everything" not in [name for name, _ in steps]
    assert "--compare-with" not in dict(steps)["Website folder"]


def test_refuses_to_overwrite_the_starting_workbook(tmp_path, capsys):
    import run_end_to_end as e2e

    book = tmp_path / "original.xlsx"
    book.write_bytes(b"PK")
    assert e2e.main([str(book), "--test-copy", str(book)]) == 2
    assert book.read_bytes() == b"PK"


def test_an_excel_automation_error_is_reported_for_retry(tmp_path, log):
    import run_end_to_end as e2e

    script = _script(tmp_path, "writer.py",
                     'import sys\nprint("Traceback (most recent call last):")\n'
                     'print("pywintypes.com_error: (-2146827864, \'OLE error 0x800a01a8\', None, None)")\n'
                     'sys.exit(1)\n')
    result = e2e.run_step("Population: town (UCL)", [sys.executable, script], log)
    assert not result["ok"] and result["excel_error"]

    other = _script(tmp_path, "other.py", 'import sys\nprint("Traceback (most recent call last):")\nsys.exit(1)\n')
    assert not e2e.run_step("Income", [sys.executable, other], log)["excel_error"]


def test_a_working_copy_that_cannot_be_written_stops_the_run(tmp_path, monkeypatch):
    import run_end_to_end as e2e

    original = tmp_path / "original.xlsx"
    original.write_bytes(b"PK")
    working = tmp_path / "test-copy.xlsx"
    (tmp_path / "~$test-copy.xlsx").write_bytes(b"")          # Excel's lock file

    def locked(src, dst):
        raise PermissionError(13, "Permission denied", str(dst))

    monkeypatch.setattr(e2e.shutil, "copyfile", locked)
    message = e2e.copy_starting_workbook(original, working)
    assert "open in Excel" in message and "test-copy.xlsx" in message

    monkeypatch.setattr(e2e, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(e2e, "build_steps", lambda *a: pytest.fail("no step may run"))
    assert e2e.main([str(original), "--test-copy", str(working)]) == 2


def test_the_working_copy_is_a_fresh_copy_of_the_starting_workbook(tmp_path):
    import run_end_to_end as e2e

    original = tmp_path / "original.xlsx"
    original.write_bytes(b"PK-new")
    working = tmp_path / "test-copy.xlsx"
    working.write_bytes(b"old contents")
    assert e2e.copy_starting_workbook(original, working) is None
    assert working.read_bytes() == b"PK-new"


def test_a_step_that_hangs_is_stopped_and_reported(tmp_path, log):
    import run_end_to_end as e2e

    script = _script(tmp_path, "hang.py", 'import time\nprint("Opening workbook", flush=True)\ntime.sleep(60)\n')
    result = e2e.run_step("Population: town (UCL)", [sys.executable, script], log, timeout_s=1)
    assert not result["ok"] and result["timed_out"]
    assert result["summary"].startswith("stopped after")
    assert result["seconds"] < 30


@pytest.mark.parametrize("path, synced", [
    ("C:/Users/me/OneDrive - The University of Queensland/Desktop/test-copy.xlsx", True),
    ("C:/Users/me/OneDrive/test-copy.xlsx", True),
    ("C:/Temp/test-copy.xlsx", False),
])
def test_a_working_copy_in_a_onedrive_folder_is_recognised(path, synced, monkeypatch):
    import run_end_to_end as e2e

    for name in ("OneDrive", "OneDriveCommercial", "OneDriveConsumer"):
        monkeypatch.delenv(name, raising=False)
    assert e2e.in_synced_folder(Path(path)) is synced


def test_the_data_year_follows_the_starting_workbook_but_never_an_unfinished_year(monkeypatch):
    import run_end_to_end as e2e

    this_year = __import__("datetime").datetime.now().year
    assert e2e.choose_data_year(this_year - 3)[0] == this_year - 2          # the year after the workbook's
    assert e2e.choose_data_year(this_year - 1)[0] == this_year - 1          # cannot add a year still under way
    assert e2e.choose_data_year(None)[0] == this_year - 1
    assert e2e.choose_data_year(2019, requested=2021) == (2021, "as requested (--year)")


def test_the_latest_year_is_read_from_the_data_sheets_headers(tmp_path):
    import openpyxl
    import run_end_to_end as e2e

    wb = openpyxl.Workbook()
    wb.active.title = "Crime"
    wb["Crime"].append([None, None, 2022, 2023, 2024])
    wb.create_sheet("Employment").append(["Smoothed"])
    wb["Employment"].append([None, None, 2023, 2024, 2025])
    wb.create_sheet("Population").append([None, None, 2040, 2050])        # projections: not read
    path = tmp_path / "book.xlsx"
    wb.save(path)
    assert e2e.latest_year_in_workbook(path) == 2025
    assert e2e.latest_year_in_workbook(tmp_path / "missing.xlsx") is None


def test_a_run_into_an_output_folder_leaves_the_workbook_and_a_plain_report(tmp_path, monkeypatch):
    import run_end_to_end as e2e

    monkeypatch.setenv("BOOMTOWN_YEAR_END", "")                          # restored after the test
    original = tmp_path / "Indicators Data-Charts 2026.xlsx"
    original.write_bytes(b"PK-workbook")
    monkeypatch.setattr(e2e, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(e2e, "SETTLE_AFTER_COPY_S", 0)
    monkeypatch.setattr(e2e, "PAUSE_BETWEEN_EXCEL_STEPS_S", 0)
    monkeypatch.setattr(e2e, "latest_year_in_workbook", lambda path: 2024)
    flagged = _script(tmp_path, "w.py", 'print("Roma Crime: FLAGGED, not written - differs from the sheet")\n'
                                        'print("Summary: 1 written, 1 flagged for review.")\n')
    monkeypatch.setattr(e2e, "build_steps", lambda args, copy: [("Crime", [sys.executable, flagged])])

    out = tmp_path / "Boomtown 2026"
    assert e2e.main([str(original), "--test-copy", str(tmp_path / "work.xlsx"), "--output-folder", str(out)]) == 0
    assert (out / "2 Workbook" / "Indicators Data-Charts 2026.xlsx").read_bytes() == b"PK-workbook"
    report = (out / "Run report.txt").read_text()
    assert "Year added        : 2025" in report
    assert "Every step finished." in report
    assert "Roma Crime: FLAGGED, not written" in report
    assert e2e.os.environ["BOOMTOWN_YEAR_END"] == "2025"


def test_report_lines_are_short_and_readable():
    import run_end_to_end as e2e

    notes = [
        "13:52:53  \x1b[33mWARNING \x1b[0m    [Shepparton] SA2 213041536 not found in SALM data",
        "13:52:53  [31mERROR   [0m    LGA file not found at C:\\cache\\salm_smoothed_lga.csv. Download it",
        "Status : ✗ FAILED",
        "[LGA] LGA file not found at C:\\cache\\salm_smoothed_lga.csv. Download it",
        "Toowoomba (West): SKIPPED (row-finding) — Could not find town 'Toowoomba (West)' in sheet 'Crime'.",
        "Yarram (Table 6): SKIPPED (row-finding) — Could not find town 'Yarram' in sheet 'Income'.",
        "Summary: 223 charts now show the new year(s); 0 not found in Excel",
        "Chinchilla weapons: WRITTEN but FLAGGED FOR REVIEW (2025 = 4.68) -> $AA$14 — SERIES: 3.9x | HISTORICAL note: 2001",
    ]
    assert e2e.report_notes(notes) == [
        "LGA file not found at C:\\cache\\salm_smoothed_lga.csv. Download it",
        "Chinchilla weapons: WRITTEN but FLAGGED FOR REVIEW (2025 = 4.68) -> $AA$14 — SERIES: 3.9x",
        "Not in this sheet, so nothing written (expected for towns the sheet does not cover): Toowoomba (West), Yarram",
    ]
