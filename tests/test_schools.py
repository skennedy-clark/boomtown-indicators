"""
tests/test_schools.py

Tests for the ACARA School Profile: summing schools by postcode,
choosing the file to download, and writing into the Education section
of the Exogenous sheet.

No network, ACARA download, workbook file or Excel is needed. The
aggregation runs on a small workbook built here with ACARA's column
headers. The writer runs against a small sheet read through
tests/fake_xlwings_sheet.py.
"""

import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parent.parent / "regional-indicators" / "transform" / "xlsx_update"))

ENROL = "Full Time Equivalent Enrolments"
STAFF = "Full Time Equivalent Teaching Staff"


# ── fetcher: aggregation ───────────────────────────────────────────────────

@pytest.fixture
def acara_file(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "DataDictionary"
    ws = wb.create_sheet("SchoolProfile 2008-2025")
    ws.append(["Calendar Year", "School Name", "State", "Postcode", STAFF, ENROL])   # column order differs from the reading order on purpose
    ws.append([2024, "Dalby State School",      "QLD", "4405", 30.0, 400.0])
    ws.append([2025, "Dalby State School",      "QLD", "4405", 31.5, 410.4])
    ws.append([2025, "Dalby State High School", "QLD", "4405", 60.2, 900.0])
    ws.append([2025, "Dalby Tiny School",       "QLD", "4405", None, None])          # no figures reported
    ws.append([2025, "Somewhere Else School",   "QLD", "4406", 99.0, 999.0])         # different postcode
    ws.append([2025, "Narrabri Public School",  "NSW", "2390", 20.0, 300.0])
    path = tmp_path / "School Profile 2008-2025.xlsx"
    wb.save(path)
    return path


def test_sums_schools_by_text_postcode_per_year(acara_file):
    from fetchers.fetch_schools import aggregate_by_postcode

    data = aggregate_by_postcode(acara_file, {"4405", "2390"})
    assert set(data) == {"4405", "2390"}                      # 4406 was not requested
    assert round(data["4405"][2025]["enrol"], 1) == 1310.4    # 410.4 + 900.0; the third school reports none
    assert round(data["4405"][2025]["staff"], 1) == 91.7
    assert len(data["4405"][2025]["schools"]) == 3            # a school with no figures still counts as a school
    assert data["4405"][2024]["enrol"] == 400.0
    assert data["2390"][2025]["enrol"] == 300.0


def test_missing_column_is_reported_not_guessed(tmp_path):
    import openpyxl
    from fetchers.fetch_schools import aggregate_by_postcode

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SchoolProfile 2025"
    ws.append(["Calendar Year", "School Name", "Postcode", ENROL])    # no teaching-staff column
    ws.append([2025, "A School", "4405", 100])
    path = tmp_path / "School Profile 2025.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="missing expected column"):
        aggregate_by_postcode(path, {"4405"})


def test_newest_year_is_tried_first_and_history_file_preferred():
    from fetchers.fetch_schools import candidate_file_names

    names = candidate_file_names(2026)
    assert names[:4] == [
        ("School Profile 2008-2026.xlsx", 2026),
        ("School Profile 2026.xlsx", 2026),
        ("School Profile 2008-2025.xlsx", 2025),
        ("School Profile 2025.xlsx", 2025),
    ]


# ── writer ─────────────────────────────────────────────────────────────────

def _sheet(tmp_path, second_label=STAFF):
    pytest.importorskip("xlwings")
    import openpyxl
    from fake_xlwings_sheet import FakeSheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Exogenous"
    for row in [
        (None, 2001, 2002, 2003),                    # sheet row 1, which is not the Education header
        ("Rainfall",),
        ("Dalby",),                                  # same town name in another section
        ("Dalby Airport", 500.0, 600.0, 700.0),
        (None,),
        ("Education", None, 2023, 2024),             # the Education section's own header row
        ("Dalby",),
        (ENROL, "FTE Enrolments", 3203.6, 3200.0),
        (second_label, "FTE Teaching Staff", 237.7, 235.7),
        (None,),
        ("Fuel", None, 2023, 2024),
        ("Dalby",),                                  # and again in the following section
        ("Average RULP Price (cents)", None, 174.4, 175.6),
    ]:
        ws.append(row)
    path = tmp_path / "exo.xlsx"
    wb.save(path)
    return FakeSheet(path, "Exogenous")


def _cache(tmp_path, town):
    path = tmp_path / f"{town.lower()}_schools.json"
    path.write_text(json.dumps({
        "town": town, "year": 2025,
        "indicators": {
            "fte_enrolments":     {"2023": 3203.6, "2024": 3200.0, "2025": 3181.4},
            "fte_teaching_staff": {"2023": 237.7,  "2024": 235.7,  "2025": 236.0},
        },
    }))
    return path


def test_writes_both_measures_into_the_education_block_only(tmp_path):
    from update_education import write_education

    sheet = _sheet(tmp_path)
    results, written, flagged = write_education(
        sheet, [_cache(tmp_path, "Dalby"), _cache(tmp_path, "Shepparton")]
    )
    assert (written, flagged) == (2, 0)
    assert sheet.writes == {
        (6, 5): 2025,       # year appended to the Education header row (E6)
        (8, 5): 3181.4,     # enrolments
        (9, 5): 236.0,      # teaching staff
    }
    assert any("No block" in line and "Shepparton" in line for line in results)


def test_block_with_unexpected_row_labels_is_refused(tmp_path):
    from update_education import write_education

    sheet = _sheet(tmp_path, second_label="Something Else")
    results, written, flagged = write_education(sheet, [_cache(tmp_path, "Dalby")])
    assert written == 0 and flagged == 2
    assert sheet.writes == {}
    assert "SKIPPED" in results[0]
