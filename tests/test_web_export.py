"""
tests/test_web_export.py -- building the website CSV folder from the
workbook (regional-indicators/transform/web_export.py).

Everything runs on a small workbook built here; no real workbook, no
Excel. The sheet shapes mirror the real ones that matter: a section
header that is also the year header, the same block name in more than
one section, a block missing its row, financial-year headers, and a
row that carries projections under future years.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "regional-indicators" / "transform"))

MAP_HEADER = "folder,file,sheet,section,block,row_label,sub_label,note\n"


@pytest.fixture
def workbook(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()

    pop = wb.active
    pop.title = "Population"
    years = list(range(2018, 2032))                      # runs on into projection years
    pop.append([None, None] + [f"{y - 1}/{str(y)[2:]}" for y in years])
    pop.append(["LGA", None] + years)                    # section header AND year header
    pop.append(["Goondiwindi", None])
    pop.append(["Population (ERP)", "Residents (LGA)"] + [10000 + i for i in range(8)] + [None] * 5 + [10999])
    pop.append(["SA2", None])
    pop.append(["Goondiwindi", None])
    pop.append(["Population (ERP)", "Residents (SA2)"] + [6000 + i for i in range(8)])
    pop.append(["Miles", None])                          # block with NO Population (ERP) row
    pop.append(["Wambo", None])
    pop.append(["Population (ERP)", "Residents (SA2)"] + [17000 + i for i in range(8)])

    inc = wb.create_sheet("Income")
    inc.append([None, None, "2017/18", "2018/19", "2019/20", "2020/21", "2021/22", "2022/23", "2023/24", "2024/25"])
    inc.append(["Dalby", None])
    inc.append([4405, None])
    inc.append(["No. wage and salary earners", None, 1, 2, 3, 4, 5, 6, 7.25, None])

    exo = wb.create_sheet("Exogenous")
    exo.append([None] + list(range(2016, 2026)))
    exo.append(["Rainfall"])
    exo.append(["Dalby "])                               # trailing space
    exo.append(["Dalby Airport 041522"] + [500.5 + i for i in range(10)])
    exo.append(["Summer"] + [300 + i for i in range(10)])

    emp = wb.create_sheet("Employment")
    emp.append([None, None] + list(range(2016, 2026)))
    emp.append(["SA2", "Label"])
    emp.append(["Wambo", "Smoothed Unemployment rate (%)"] + [3.35] * 10)

    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def _map(tmp_path, *lines):
    path = tmp_path / "map.csv"
    path.write_text(MAP_HEADER + "\n".join(lines) + "\n", encoding="utf-8")
    return path


def _read(path):
    return path.read_bytes().decode("utf-8")


# ── end to end ─────────────────────────────────────────────────────────────

def test_export_writes_one_csv_per_map_line_in_town_folders(workbook, tmp_path):
    from web_export import export

    map_path = _map(
        tmp_path,
        "Goondiwindi,Population - LGA,Population,LGA,Goondiwindi,Population (ERP),,",
        "Goondiwindi,Population - district,Population,SA2,Goondiwindi,Population (ERP),,",
        "Dalby (Wambo),Number of earners,Income,,Dalby,No. wage and salary earners,,",
        "Dalby (Wambo),Environment - total rainfall,Exogenous,Rainfall,Dalby,@next,,",
        "Dalby (Wambo),Unemployment rate,Employment,SA2,,Wambo,,",
    )
    out = tmp_path / "3 Web Content"
    report, problems = export(workbook, out, map_path, last_year=2026)
    assert problems == 0
    assert sorted(p.relative_to(out).as_posix() for p in out.glob("*/*.csv")) == [
        "Dalby (Wambo)/Environment - total rainfall.csv",
        "Dalby (Wambo)/Number of earners.csv",
        "Dalby (Wambo)/Unemployment rate.csv",
        "Goondiwindi/Population - LGA.csv",
        "Goondiwindi/Population - district.csv",
    ]

    lga = _read(out / "Goondiwindi" / "Population - LGA.csv")
    header, data, tail = lga.split("\r\n")
    assert tail == ""                                           # ends with CRLF, two lines only
    assert header.split(",") == [str(y) for y in range(2000, 2026)]   # same header in every file
    cells = data.split(",")
    assert cells[0] == "" and cells[18:] == [str(10000 + i) for i in range(8)]   # 2018-2025
    assert "10999" not in lga                                   # the 2031 projection is left out
    assert any("2031" in line and "Population - LGA" in line for line in report)

    # same block name, different section, different numbers
    assert _read(out / "Goondiwindi" / "Population - district.csv").split("\r\n")[1].endswith("6007")

    # financial years count as the year they end in: 2023/24 -> 2024
    earners = _read(out / "Dalby (Wambo)" / "Number of earners.csv").split("\r\n")[1].split(",")
    assert earners[24] == "7.25" and earners[25] == ""          # 2024 has it, 2025 blank

    # "@next" takes the row straight after the block heading, whatever its label
    rain = _read(out / "Dalby (Wambo)" / "Environment - total rainfall.csv").split("\r\n")[1].split(",")
    assert rain[16] == "500.5" and rain[25] == "509.5"


def test_missing_row_is_reported_not_taken_from_the_next_block(workbook, tmp_path):
    from web_export import export

    map_path = _map(
        tmp_path,
        "Miles,Population - district,Population,SA2,Miles,Population (ERP),,",
        "Dalby (Wambo),Population - district,Population,SA2,Wambo,Population (ERP),,",
    )
    out = tmp_path / "out"
    report, problems = export(workbook, out, map_path, last_year=2026)
    assert problems == 1
    assert not (out / "Miles").exists()                         # Wambo's row was NOT used for Miles
    assert (out / "Dalby (Wambo)" / "Population - district.csv").exists()
    assert any("Miles/Population - district.csv" in line and "no row labelled" in line for line in report)


def test_two_map_lines_for_the_same_file_are_refused(tmp_path):
    from web_export import load_map

    map_path = _map(
        tmp_path,
        "Dalby,Theft,Crime,,Dalby,Other Theft,,",
        "Dalby,Theft,Crime,,Dalby,Theft,,",
    )
    with pytest.raises(ValueError, match="both produce"):
        load_map(map_path)


def test_dry_run_writes_nothing(workbook, tmp_path):
    from web_export import export

    map_path = _map(tmp_path, "Dalby (Wambo),Unemployment rate,Employment,SA2,,Wambo,,")
    out = tmp_path / "out"
    report, problems = export(workbook, out, map_path, dry_run=True, last_year=2026)
    assert problems == 0 and not out.exists()
    assert any(line.startswith("Would write") for line in report)


# ── numbers ────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("value, text", [
    (6565, "6565"),
    (6565.0, "6565"),
    (62.245401, "62.2454"),
    (86.999184, "86.99918"),
    (440381.25, "440381.3"),          # half-up, as Excel shows it
    (0.6821282, "0.682128"),
    (0.07429973, "0.0743"),
    (1.6749999999999998, "1.675"),
    (0, "0"),
    (None, ""),
])
def test_numbers_are_written_as_excel_showed_them(value, text):
    from web_export import format_number

    assert format_number(value) == text


# ── comparing with an earlier export ───────────────────────────────────────

def test_reads_earlier_files_however_they_were_saved(tmp_path):
    from web_export import read_existing_csv

    quoted = tmp_path / "quoted.csv"
    quoted.write_bytes(b'"2000","2001","2002"\r\n"","62.2454","60.2449"')
    assert read_existing_csv(quoted) == {2001: 62.2454, 2002: 60.2449}

    by_hand = tmp_path / "by_hand.csv"
    by_hand.write_bytes(b'2000,2001,2002\r\n," 10,141 ","10,126"\r\n,,\r\n')
    assert read_existing_csv(by_hand) == {2001: 10141.0, 2002: 10126.0}


def test_compare_separates_new_year_from_changed_history(workbook, tmp_path):
    from web_export import export

    map_path = _map(
        tmp_path,
        "Goondiwindi,Population - LGA,Population,LGA,Goondiwindi,Population (ERP),,",
        "Goondiwindi,Population - district,Population,SA2,Goondiwindi,Population (ERP),,",
    )
    old = tmp_path / "old" / "Goondiwindi"
    old.mkdir(parents=True)
    hdr = ",".join(str(y) for y in range(2000, 2025))
    (old / "Population - LGA.csv").write_text(      # same history, one year shorter
        hdr + "\r\n" + ",".join([""] * 18 + [str(10000 + i) for i in range(7)]) + "\r\n")
    (old / "Population - district.csv").write_text( # 2019 was different
        hdr + "\r\n" + ",".join([""] * 18 + ["6000", "5990"] + [str(6002 + i) for i in range(5)]) + "\r\n")

    report, _ = export(workbook, tmp_path / "out", map_path, compare_with=tmp_path / "old",
                       dry_run=True, last_year=2026)
    text = "\n".join(report)
    assert "0 identical, 1 same history plus new year(s), 1 with differences" in text
    assert "DIFF  Goondiwindi/Population - district.csv" in text and "2019: 5990 -> 6001" in text


def test_comparing_with_the_folder_being_written_reads_it_first(workbook, tmp_path):
    """Last year's folder is often the folder being regenerated. The
    comparison must describe it as it WAS, not as just rewritten."""
    from web_export import export

    map_path = _map(tmp_path, "Goondiwindi,Population - district,Population,SA2,Goondiwindi,Population (ERP),,")
    out = tmp_path / "3 Web Content"
    (out / "Goondiwindi").mkdir(parents=True)
    hdr = ",".join(str(y) for y in range(2000, 2025))
    (out / "Goondiwindi" / "Population - district.csv").write_text(
        hdr + "\r\n" + ",".join([""] * 18 + ["6000", "5990"] + [str(6002 + i) for i in range(5)]) + "\r\n")

    report, _ = export(workbook, out, map_path, compare_with=out, last_year=2026)
    assert "1 with differences" in "\n".join(report)                    # saw the OLD 2019 figure
    assert "6001" in _read(out / "Goondiwindi" / "Population - district.csv")   # and then replaced it